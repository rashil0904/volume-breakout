# -*- coding: utf-8 -*-
"""pead_final_report.py — consolidated Excel report for the 3 finalized PEAD books, flat Rs.50,000 per
trade (independent sizing). Books: FNO SHORT (K4,T+3,close-SL@T0high,8% tgt); FNO LONG (K4,T+15,
close-SL@T0low,no tgt); NON-F&O LONG (K7,T+10,close-SL@T0low,17% tgt). Guards applied (stale prev-VWAP,
(symbol,T0) de-dup, |reaction|>50% corp-action). Returns net of 0.20% round-trip. Sheets: Summary,
Quarterly, and per-book trade lists. -> pead_final_report.xlsx
"""
import sys
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import run_backtest as rb

ANN = rb.RESULTS / "bse_results_announcements" / "bse_results_announcements_combined.csv"
FO_CSV = rb.RESULTS / "bse_results_announcements" / "bse_results_announcements.csv"
UNIV = rb.BASE / "data" / "nifty500_universe.csv"
DAILY = rb.BASE / "data" / "daily_ohlcv_all.parquet"
VWAPD = rb.BASE / "data" / "daily_ohlcv_vwapclose.parquet"
OUT = rb.RESULTS / "pead_final_report.xlsx"
CAP, COST = 50000.0, 0.20
IS_YEARS = {2022, 2023, 2024}


def load():
    a = pd.read_csv(ANN, dtype=str)
    a["ann_date"] = pd.to_datetime(a["announcement_date"], format="%d-%m-%Y", errors="coerce").dt.date
    a["dtm"] = pd.to_datetime(a["announcement_datetime"], format="%d-%m-%Y %H:%M", errors="coerce")
    a = a.dropna(subset=["ann_date", "dtm"]).sort_values("dtm").drop_duplicates(["symbol", "quarter"], keep="first")
    d = pd.read_parquet(DAILY, columns=["symbol", "date", "open", "high", "low", "close"]); d["date"] = pd.to_datetime(d["date"]).dt.date
    vw = pd.read_parquet(VWAPD, columns=["symbol", "date", "close"]); vw["date"] = pd.to_datetime(vw["date"]).dt.date
    vmap = {(s, dt): c for s, dt, c in zip(vw["symbol"], vw["date"], vw["close"])}
    sd = {}
    for sym, g in d.groupby("symbol", sort=False):
        g = g.sort_values("date")
        sd[sym] = (g["date"].values, g["open"].values.astype(float), g["high"].values.astype(float),
                   g["low"].values.astype(float), g["close"].values.astype(float), {dt: i for i, dt in enumerate(g["date"].values)})
    return a, sd, vmap


def fy_quarter(dt):
    m, y = dt.month, dt.year
    if m >= 4:
        fy = y + 1; q = 1 if m <= 6 else 2 if m <= 9 else 3
    else:
        fy = y; q = 4
    return f"FY{fy % 100:02d} Q{q}", fy * 10 + q


def simulate(direction, entry, sl, n, tgt_pct, fo, fh, fl, fc):
    if direction == "long":
        tgt = entry * (1 + tgt_pct / 100.0) if tgt_pct else None
        for j in range(n):
            if tgt is not None and fh[j] >= tgt:
                return max(tgt, fo[j]), "target", j + 1
            if fc[j] <= sl:
                return fc[j], "SL", j + 1
        return fc[n - 1], "holding-end", n
    tgt = entry * (1 - tgt_pct / 100.0) if tgt_pct else None
    for j in range(n):
        if tgt is not None and fl[j] <= tgt:
            return min(tgt, fo[j]), "target", j + 1
        if fc[j] >= sl:
            return fc[j], "SL", j + 1
    return fc[n - 1], "holding-end", n


def build_book(a, sd, vmap, syms, direction, K, n, tgt_pct):
    rows = []; seen = set()
    for r in a.itertuples():
        if r.symbol not in syms:
            continue
        s = sd.get(r.symbol)
        if s is None:
            continue
        dates, op, hi, lo, cl, dmap = s
        ps = dmap.get(r.ann_date); pn = int(np.searchsorted(dates, r.ann_date, "right")); pn = pn if pn < len(dates) else None
        p = pn if r.announcement_session in ("post_market", "during_market") else (ps if ps is not None else pn)
        if p is None or p < 1:
            continue
        if (dates[p] - dates[p - 1]).days > 6:                 # GUARD1 stale prev-VWAP
            continue
        key = (r.symbol, int(p))                                # GUARD2 de-dup (symbol,T0)
        if key in seen:
            continue
        seen.add(key)
        entry = cl[p]; pv = vmap.get((r.symbol, dates[p - 1]), np.nan)
        if not (entry > 0 and pv == pv and pv > 0):
            continue
        rr = (entry - pv) / pv * 100.0
        if abs(rr) > 50.0:                                      # GUARD3 corp-action
            continue
        if direction == "long" and rr < K:
            continue
        if direction == "short" and rr > -K:
            continue
        if p + n >= len(dates):
            continue
        sl = lo[p] if direction == "long" else hi[p]
        fo, fh, fl, fc = op[p + 1:p + 1 + n], hi[p + 1:p + 1 + n], lo[p + 1:p + 1 + n], cl[p + 1:p + 1 + n]
        exit_px, reason, dh = simulate(direction, entry, sl, n, tgt_pct, fo, fh, fl, fc)
        gross = (exit_px - entry) / entry * 100.0 if direction == "long" else (entry - exit_px) / entry * 100.0
        net = gross - COST; inr = net / 100.0 * CAP
        fq, fqk = fy_quarter(dates[p])
        rows.append({"symbol": r.symbol, "announcement_date": r.announcement_date, "announcement_time": r.announcement_time,
                     "session": r.announcement_session, "T0_date": str(dates[p]), "entry_price": round(entry, 2),
                     "exit_date": str(dates[p + dh]), "exit_price": round(exit_px, 2), "exit_reason": reason,
                     "return_pct_net": round(net, 3), "return_inr": round(inr, 0), "holding_days": dh,
                     "year": int(str(dates[p])[:4]), "fy_quarter": fq, "fqk": fqk})
    return pd.DataFrame(rows)


def metrics(df):
    if df.empty:
        return {k: 0 for k in range(13)}
    inr = df["return_inr"]; win = inr > 0; loss = inr < 0
    return {"Total trades": len(df), "Winning trades": int(win.sum()), "Losing trades": int(loss.sum()),
            "Trades hit SL": int((df.exit_reason == "SL").sum()), "Win rate (%)": round(win.mean() * 100, 1),
            "Avg return / trade (INR)": round(inr.mean(), 0), "Median return / trade (INR)": round(inr.median(), 0),
            "Avg holding (days, realized)": round(df.holding_days.mean(), 1), "Total PnL (INR)": round(inr.sum(), 0),
            "Avg positive PnL (INR)": round(inr[win].mean(), 0) if win.any() else 0,
            "Avg negative PnL (INR)": round(inr[loss].mean(), 0) if loss.any() else 0,
            "Max profit (INR)": round(inr.max(), 0), "Max loss (INR)": round(inr.min(), 0),
            "-- In-sample PnL (T0 2022-24)": round(inr[df.year.isin(IS_YEARS)].sum(), 0),
            "-- Out-of-sample PnL (T0 2025+)": round(inr[~df.year.isin(IS_YEARS)].sum(), 0)}


# ---------- Excel styling ----------
HDR = PatternFill("solid", fgColor="1F4E78"); SUB = PatternFill("solid", fgColor="D9E1F2")
WHITE = Font(bold=True, color="FFFFFF"); BOLD = Font(bold=True)
RIGHT = Alignment(horizontal="right"); THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
RUP = '"₹"#,##0'


def sheet_summary(ws, books):
    ws["A1"] = "PEAD — Consolidated Results (3 finalized books)"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Sizing: flat ₹50,000 per trade, independent (not a shared pool).  Returns NET of 0.20% round-trip cost.  "
                "Guards: stale-VWAP/(symbol,T0) de-dup/|reaction|>50% corp-action removed.  OOS = T0 calendar 2025+ (validated).")
    ws["A2"].font = Font(italic=True, size=9)
    order = [("TOTAL (ALL THREE BOOKS)", "TOTAL"), ("FNO SHORT  (K4 · T+3 · close-SL@T0 high · 8% target)", "FNO Short"),
             ("FNO LONG  (K4 · T+15 · close-SL@T0 low · no target)", "FNO Long"),
             ("NON-F&O LONG  (K7 · T+10 · close-SL@T0 low · 17% target)", "Non-F&O Long")]
    row = 4
    for title, kb in order:
        ws.cell(row, 1, title).font = WHITE
        for c in (1, 2):
            ws.cell(row, c).fill = HDR
        ws.cell(row, 2, "value").font = WHITE
        row += 1
        for k, v in metrics(books[kb]).items():
            ws.cell(row, 1, k).font = BOLD if k.startswith("--") else Font()
            cell = ws.cell(row, 2, v); cell.alignment = RIGHT
            if "INR" in k or "PnL" in k or k.startswith("--"):
                cell.number_format = RUP
            elif "%" in k:
                cell.number_format = '0.0'
            for c in (1, 2):
                ws.cell(row, c).border = THIN
            row += 1
        row += 1                                               # spacer
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 18


def sheet_quarterly(ws, allt):
    books = ["FNO Short", "FNO Long", "Non-F&O Long", "Total"]
    mets = ["trades", "win%", "PnL(₹)", "avg/trade(₹)"]
    ws.cell(1, 1, "QUARTERLY SUMMARY — Indian FY (Q1 Apr-Jun … Q4 Jan-Mar), by T0 date").font = Font(bold=True, size=12)
    # header rows (2 levels)
    ws.cell(3, 1, "FY-Quarter").font = WHITE; ws.cell(3, 1).fill = HDR
    col = 2
    span = {}
    for b in books:
        span[b] = col
        ws.cell(3, col, b).font = WHITE
        for i in range(4):
            ws.cell(3, col + i).fill = HDR; ws.cell(3, col + i).font = WHITE
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 3)
        ws.cell(3, col).alignment = Alignment(horizontal="center")
        for i, mm in enumerate(mets):
            ws.cell(4, col + i, mm).font = BOLD; ws.cell(4, col + i).fill = SUB; ws.cell(4, col + i).alignment = RIGHT
        col += 4
    ws.cell(4, 1, "").fill = SUB
    fqs = allt.sort_values("fqk")["fy_quarter"].drop_duplicates().tolist()
    r = 5
    for fq in fqs:
        ws.cell(r, 1, fq).font = BOLD
        for b in books:
            sub = allt if b == "Total" else allt[allt.book == b]
            sub = sub[sub.fy_quarter == fq]
            c0 = span[b]
            if len(sub):
                vals = [len(sub), round((sub.return_inr > 0).mean() * 100, 1), round(sub.return_inr.sum(), 0), round(sub.return_inr.mean(), 0)]
            else:
                vals = [0, 0, 0, 0]
            for i, v in enumerate(vals):
                cell = ws.cell(r, c0 + i, v); cell.alignment = RIGHT
                if i == 1:
                    cell.number_format = '0.0'
                elif i >= 2:
                    cell.number_format = RUP
        r += 1
    # totals row
    ws.cell(r, 1, "ALL QUARTERS").font = Font(bold=True, color="1F4E78")
    for b in books:
        sub = allt if b == "Total" else allt[allt.book == b]
        c0 = span[b]
        vals = [len(sub), round((sub.return_inr > 0).mean() * 100, 1), round(sub.return_inr.sum(), 0), round(sub.return_inr.mean(), 0)] if len(sub) else [0, 0, 0, 0]
        for i, v in enumerate(vals):
            cell = ws.cell(r, c0 + i, v); cell.font = BOLD; cell.alignment = RIGHT
            cell.number_format = '0.0' if i == 1 else (RUP if i >= 2 else 'General')
    ws.column_dimensions["A"].width = 12
    for c in range(2, 2 + 16):
        ws.column_dimensions[get_column_letter(c)].width = 12


TCOLS = ["symbol", "announcement_date", "announcement_time", "session", "T0_date", "entry_price",
         "exit_date", "exit_price", "exit_reason", "return_pct_net", "return_inr", "holding_days"]


def sheet_trades(ws, df, title):
    ws.cell(1, 1, title).font = Font(bold=True, size=12)
    for j, c in enumerate(TCOLS, 1):
        cell = ws.cell(3, j, c); cell.font = WHITE; cell.fill = HDR
    for i, rr in enumerate(df.sort_values("T0_date")[TCOLS].itertuples(index=False), 4):
        for j, v in enumerate(rr, 1):
            cell = ws.cell(i, j, v)
            if TCOLS[j - 1] == "return_inr":
                cell.number_format = RUP
            elif TCOLS[j - 1] == "return_pct_net":
                cell.number_format = '0.00'
    for j, c in enumerate(TCOLS, 1):
        ws.column_dimensions[get_column_letter(j)].width = 16 if "date" in c or c == "symbol" else 13


def main():
    a, sd, vmap = load()
    fno = set(pd.read_csv(FO_CSV, usecols=["symbol"])["symbol"].unique())
    u = pd.read_csv(UNIV); nonfo = set(u[~u["fno_eligible"]]["symbol"])
    books = {
        "FNO Short": build_book(a, sd, vmap, fno, "short", 4.0, 3, 8.0),
        "FNO Long": build_book(a, sd, vmap, fno, "long", 4.0, 15, None),
        "Non-F&O Long": build_book(a, sd, vmap, nonfo, "long", 7.0, 10, 17.0),
    }
    allt = pd.concat([df.assign(book=b) for b, df in books.items()], ignore_index=True)
    books["TOTAL"] = allt
    for b, df in books.items():
        if b != "TOTAL":
            print(f"{b:14s}: {len(df):4d} trades | PnL ₹{df.return_inr.sum():,.0f} | win {round((df.return_inr>0).mean()*100,1)}%")
    print(f"{'TOTAL':14s}: {len(allt):4d} trades | PnL ₹{allt.return_inr.sum():,.0f}")

    wb = Workbook()
    sheet_summary(wb.active, books); wb.active.title = "Summary"
    sheet_quarterly(wb.create_sheet("Quarterly_Summary"), allt)
    sheet_trades(wb.create_sheet("FNO_Short_Trades"), books["FNO Short"], "FNO SHORT — all trades (K4·T+3·close-SL·8% tgt)")
    sheet_trades(wb.create_sheet("FNO_Long_Trades"), books["FNO Long"], "FNO LONG — all trades (K4·T+15·close-SL·no tgt)")
    sheet_trades(wb.create_sheet("NonFNO_Long_Trades"), books["Non-F&O Long"], "NON-F&O LONG — all trades (K7·T+10·close-SL·17% tgt)")
    wb.save(OUT)
    print(f"\nSaved -> {OUT}")


if __name__ == "__main__":
    main()
