# -*- coding: utf-8 -*-
"""
short_target_cover_sweep.py — 2-D short-exit sweep (cover time × short target %) on the FINALIZED
long config (conditional split t1=09:25, t2=11:59, long target 17%). Composite strategy: Category C
entry 15:21, A/B unchanged, LONG side fixed — only the double-down short's exit varies.

SHORT EXIT per trade, per (cover_time T in 14:30-15:00, short target X in 3-20% or no-target):
  short opens at the trade's LONG EXIT price/time (filled qty).
  target_level = short_open × (1 − X/100)   (price falls X% → short +X% in profit).
  scan 1-min LOWS from AFTER the long-exit minute up to (not incl.) T:
    low <= target_level  → cover at target_level (limit-fill, "short_target_hit"); first candle wins.
    never hit by T       → cover at T's 1-min open ("short_cover_at_time").
  no-target: always cover at T.
  short_pnl = shares × (short_open − cover);  combined = long_pnl + short_pnl.
Costs gross / net_A (long0.23%+short0.10%) / net_B (long0.38%+short0.10%); short notional = shares×short_open
(fixed per trade), so short cost is cover-independent. 31 covers × 19 targets = 589 combos.

Flags: (a) short target limit-fill at short_open×(1−X/100) on a 1-min LOW; (b) scan long-exit→cover,
else cover at time; (c) 3-20% + no-target = 19 × 31 = 589; (d) short_open = long exit price;
(e) combined pnl, short leg isolated + by-origin (17% long-target vs 09:25 vs 11:59 exits).
"""
import sys, time
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import run_backtest as rb
import uc_staggered_dd_report as R

OUTDIR = rb.RESULTS / "short_target_cover_sweep"
IST = R.IST
BASE_POOL, R023, R038, SR = R.BASE_POOL, R.LONG_023, R.LONG_038, R.SHORT_RATE
HM_C_ENTRY = 921
LONG_T1, LONG_T2, LONG_TGT = 565, 719, 17            # 09:25 / 11:59 / 17% long
GRID_HMS = np.arange(560, 721)                        # 09:20..12:00 (trade-set consistency filter)
LHIGH_HMS = np.arange(555, 720)                       # long-target scan (09:15..11:59)
LOW_HMS = np.arange(555, 901)                         # short-target scan (09:15..15:00)
COVER_HMS = np.arange(870, 901)                       # 14:30..15:00 (31)
SHORT_TGTS = list(range(3, 21)) + [None]              # 3..20 + no-target (19)
INCUMBENT = ("15:00", "no_target")                    # fixed 3pm cover, no short target
KEY = "net_A_total_return_fixedbase_pct"


def lbl(hm):
    return f"{int(hm)//60:02d}:{int(hm)%60:02d}"


def build_cache():
    diag = pd.read_csv(rb.RESULTS / "diagnostic_table.csv",
                       usecols=["symbol", "date", "passes_all_three",
                                "prev_day_vwap_close", "entry_price_315pm"], parse_dates=["date"])
    diag["date"] = diag["date"].dt.date
    Q = diag[diag["passes_all_three"] == True]
    print("Scanning engine …", flush=True)
    records = []
    for s in sorted(Q["symbol"].unique()):
        records += R.scan_symbol(s, Q[Q["symbol"] == s])
    T_def, day_diag = R.size_and_price(records)
    per_C = {d: v["per_C"] for d, v in day_diag.items()}
    trades = []
    for _, r in T_def[T_def["category"].isin(["A", "B"])].iterrows():
        trades.append(dict(symbol=r["symbol"], entry_date=r["entry_date"], exit_date=r["exit_date"],
                           entry_price=r["avg_entry"], shares=float(r["shares"]),
                           cap=float(r["capital_deployed"]), per_C=None))
    for r in records:
        if r["category"] == "C" and r["entered"] and per_C.get(r["entry_date"], 0) > 0:
            trades.append(dict(symbol=r["symbol"], entry_date=r["entry_date"], exit_date=r["next_date"],
                               entry_price=None, shares=None, cap=None, per_C=per_C[r["entry_date"]]))
    by_sym = {}
    for t in trades:
        by_sym.setdefault(t["symbol"], []).append(t)
    rows, dropped = [], 0
    t0 = time.time()
    for si, (sym, ts) in enumerate(by_sym.items(), 1):
        pq = rb.MASTER_DIR / f"{sym}.parquet"
        if not pq.exists():
            dropped += len(ts); continue
        raw = pd.read_parquet(pq)
        tstamp = pd.to_datetime(raw["timestamp"], utc=True).dt.tz_convert(IST)
        raw = raw.assign(date=tstamp.dt.date, hm=tstamp.dt.hour * 60 + tstamp.dt.minute)
        by_date = {d: g for d, g in raw.groupby("date")}
        for t in ts:
            ed, xd = t["entry_date"], t["exit_date"]
            if xd is None or xd not in by_date:
                dropped += 1; continue
            xg = by_date[xd]
            xo = dict(zip(xg["hm"].values, xg["open"].values.astype(float)))
            xh = dict(zip(xg["hm"].values, xg["high"].values.astype(float)))
            xl = dict(zip(xg["hm"].values, xg["low"].values.astype(float)))
            grid = np.array([xo.get(h, np.nan) for h in GRID_HMS])
            lhigh = np.array([xh.get(h, np.nan) for h in LHIGH_HMS])
            low = np.array([xl.get(h, np.nan) for h in LOW_HMS])
            cov = np.array([xo.get(h, np.nan) for h in COVER_HMS])
            if t["per_C"] is not None:
                eg = by_date.get(ed)
                p = dict(zip(eg["hm"].values, eg["open"].values.astype(float))).get(HM_C_ENTRY, np.nan) if eg is not None else np.nan
                if not (p == p and p > 0):
                    dropped += 1; continue
                shares = float(np.floor(t["per_C"] / p))
                if shares <= 0:
                    dropped += 1; continue
                entry_price, cap = p, shares * p
            else:
                entry_price, shares, cap = t["entry_price"], t["shares"], t["cap"]
            if np.isnan(grid).any() or np.isnan(cov).any():          # same trade set as long sweep + cover grid
                dropped += 1; continue
            rows.append((entry_price, shares, cap, grid, lhigh, low, cov))
        if si % 300 == 0:
            print(f"  …{si}/{len(by_sym)} ({time.time()-t0:.0f}s), kept {len(rows):,}", flush=True)
    C = dict(entry=np.array([r[0] for r in rows], float), shares=np.array([r[1] for r in rows], float),
             cap=np.array([r[2] for r in rows], float),
             grid=np.vstack([r[3] for r in rows]), lhigh=np.vstack([r[4] for r in rows]),
             low=np.vstack([r[5] for r in rows]), cov=np.vstack([r[6] for r in rows]))
    print(f"  kept {len(rows):,} trades; dropped {dropped:,}.", flush=True)
    return C


def long_side(C):
    """Per-trade long exit price + time + long_pnl + short-origin, for the fixed 09:25/11:59/17% split."""
    entry = C["entry"]
    lt = entry * (1 + LONG_TGT / 100.0)
    crossed = C["lhigh"] >= lt[:, None]
    ever = crossed.any(axis=1); first = np.argmax(crossed, axis=1)
    lt_hm = np.where(ever, LHIGH_HMS[first], 10 ** 9)                # long-target hit minute
    i1 = int(np.where(GRID_HMS == LONG_T1)[0][0]); i2 = int(np.where(GRID_HMS == LONG_T2)[0][0])
    o1, o2 = C["grid"][:, i1], C["grid"][:, i2]
    pos = o1 > entry
    le_hm = np.full(len(entry), LONG_T2); le_px = o2.copy(); origin = np.array(["t2_1159"] * len(entry), dtype=object)
    # non-pos with target in t1..t2
    m = (~pos) & (lt_hm <= LONG_T2)
    le_hm = np.where(m, lt_hm, le_hm); le_px = np.where(m, lt, le_px); origin[m] = "long_target"
    # positive at t1 -> exit t1
    le_hm = np.where(pos, LONG_T1, le_hm); le_px = np.where(pos, o1, le_px); origin[pos] = "t1_0925"
    # target before t1 overrides (pos or not)
    m = lt_hm <= LONG_T1
    le_hm = np.where(m, lt_hm, le_hm); le_px = np.where(m, lt, le_px); origin[m] = "long_target"
    long_pnl = C["shares"] * (le_px - entry)
    return le_hm, le_px, long_pnl, origin


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    C = build_cache()
    N = len(C["entry"])
    le_hm, short_open, long_pnl, origin = long_side(C)
    shares, cap = C["shares"], C["cap"]
    long_cost_A, long_cost_B = R023 * cap, R038 * cap
    short_notl = shares * short_open
    short_cost = SR * short_notl                                    # fixed per trade
    # lows valid only AFTER the long exit minute
    valid = LOW_HMS[None, :] > le_hm[:, None]
    lows_masked = np.where(valid, C["low"], np.inf)
    covers = C["cov"]                                               # N × 31 (cover opens)

    rows, mat = [], {}
    t0 = time.time()
    for X in SHORT_TGTS:
        xlabel = "no_target" if X is None else f"{X}%"
        if X is None:
            short_hit_hm = np.full(N, 10 ** 9)
        else:
            level = short_open * (1 - X / 100.0)
            cr = lows_masked <= level[:, None]
            ever = cr.any(axis=1); first = np.argmax(cr, axis=1)
            short_hit_hm = np.where(ever, LOW_HMS[first], 10 ** 9)
        for ki, tc in enumerate(COVER_HMS):
            if X is None:
                hit = np.zeros(N, bool); cover_px = covers[:, ki]
            else:
                hit = short_hit_hm < tc
                cover_px = np.where(hit, level, covers[:, ki])
            short_pnl = shares * (short_open - cover_px)
            comb = long_pnl + short_pnl
            netA = comb - long_cost_A - short_cost
            netB = comb - long_cost_B - short_cost
            row = {"cover_time": lbl(int(tc)), "short_target": xlabel, "n_trades": N,
                   "short_leg_total_pnl_inr": round(float(short_pnl.sum()), 0),
                   "short_leg_return_fixedbase_pct": round(float(short_pnl.sum() / BASE_POOL * 100), 4),
                   "short_win_rate_pct": round(float((short_pnl > 0).mean() * 100), 2),
                   "avg_short_return_pct": round(float(((short_open - cover_px) / short_open * 100).mean()), 4),
                   "pct_shorts_hitting_target": round(float(hit.mean() * 100), 2)}
            for tag, p in [("gross", comb), ("net_A", netA), ("net_B", netB)]:
                r = p / cap * 100
                row[f"{tag}_total_return_fixedbase_pct"] = round(float(p.sum() / BASE_POOL * 100), 4)
                row[f"{tag}_total_pnl_inr"] = round(float(p.sum()), 0)
                row[f"{tag}_win_rate_pct"] = round(float((p > 0).mean() * 100), 2)
                row[f"{tag}_avg_return_per_trade_pct"] = round(float(r.mean()), 4)
                row[f"{tag}_median_return_per_trade_pct"] = round(float(np.median(r)), 4)
            rows.append(row)
            mat[(xlabel, lbl(int(tc)))] = row[KEY]
    df = pd.DataFrame(rows).sort_values(KEY, ascending=False).reset_index(drop=True)
    df.insert(0, "rank", range(1, len(df) + 1))
    print(f"combos: {len(df)} ({time.time()-t0:.0f}s)")

    best = df.iloc[0]
    inc = df[(df["cover_time"] == INCUMBENT[0]) & (df["short_target"] == INCUMBENT[1])].iloc[0]

    # target vs no-target
    with_t = df[df["short_target"] != "no_target"].iloc[0]
    no_t = df[df["short_target"] == "no_target"].iloc[0]

    # pct hitting target across X (at the best cover time)
    bc = best["cover_time"]
    xorder = [f"{x}%" for x in range(3, 21)] + ["no_target"]
    pct_tbl = (df[df["cover_time"] == bc].set_index("short_target")
               .reindex(xorder)[["pct_shorts_hitting_target", KEY]].reset_index())

    # by-origin: best (cover,target) by that origin's short-leg total pnl
    by_org = []
    for org in ["long_target", "t1_0925", "t2_1159"]:
        msk = origin == org
        sub_shares, sub_so, sub_cap = shares[msk], short_open[msk], cap[msk]
        sub_valid = LOW_HMS[None, :] > le_hm[msk][:, None]
        sub_lows = np.where(sub_valid, C["low"][msk], np.inf)
        sub_cov = covers[msk]
        best_o = None
        for X in SHORT_TGTS:
            if X is None:
                sh = np.full(msk.sum(), 10 ** 9); lv = None
            else:
                lv = sub_so * (1 - X / 100.0); cr = sub_lows <= lv[:, None]
                ev = cr.any(axis=1); fi = np.argmax(cr, axis=1); sh = np.where(ev, LOW_HMS[fi], 10 ** 9)
            for ki, tc in enumerate(COVER_HMS):
                cp = sub_cov[:, ki] if X is None else np.where(sh < tc, lv, sub_cov[:, ki])
                spnl = (sub_shares * (sub_so - cp)).sum()
                if best_o is None or spnl > best_o[2]:
                    best_o = (lbl(int(tc)), "no_target" if X is None else f"{X}%", spnl)
        by_org.append({"origin": org, "n": int(msk.sum()), "best_cover": best_o[0],
                       "best_short_target": best_o[1], "short_leg_pnl_inr": round(float(best_o[2]), 0)})
    by_org = pd.DataFrame(by_org)

    # ── save Excel ──
    def sty(ws):
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2E74B5")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 22)
    with pd.ExcelWriter(OUTDIR / "short_target_cover_sweep.xlsx", engine="openpyxl") as w:
        df.to_excel(w, sheet_name="all_589", index=False)
        df.head(20).to_excel(w, sheet_name="top20", index=False)
        pct_tbl.to_excel(w, sheet_name="pct_hit_target_by_X", index=False)
        by_org.to_excel(w, sheet_name="by_origin", index=False)
        pd.DataFrame([best, inc], index=["best", "incumbent_3pm_notarget"]).to_excel(w, sheet_name="best_vs_incumbent")
        for ws in w.sheets.values():
            sty(ws)

    # ── heatmap ──
    covs = [lbl(int(h)) for h in COVER_HMS]
    M = np.array([[mat[(xl, cv)] for cv in covs] for xl in xorder])
    fig, ax = plt.subplots(figsize=(11, 7))
    im = ax.imshow(M, aspect="auto", cmap="RdYlGn", origin="upper")
    ax.set_xticks(range(len(covs))); ax.set_xticklabels(covs, rotation=90, fontsize=7)
    ax.set_yticks(range(len(xorder))); ax.set_yticklabels(xorder, fontsize=8)
    ax.set_xlabel("short cover time"); ax.set_ylabel("short target %")
    ax.set_title("Combined net_A total_return %  (short cover × short target)  — long fixed 09:25/11:59/17%",
                 fontweight="bold", fontsize=11)
    fig.colorbar(im, ax=ax, shrink=0.8, label="net_A total_return %")
    bi = xorder.index(best["short_target"]); bj = covs.index(best["cover_time"])
    ax.scatter([bj], [bi], marker="*", s=200, c="black")
    fig.tight_layout(); fig.savefig(OUTDIR / "short_cover_x_target_heatmap.png", dpi=130); plt.close(fig)

    # ── two 1D cuts ──
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(14, 5))
    cut1 = df[df["short_target"] == best["short_target"]].copy()
    cut1["cm"] = cut1["cover_time"].map(lambda s: int(s[:2]) * 60 + int(s[3:])); cut1 = cut1.sort_values("cm")
    a1.plot(cut1["cover_time"], cut1[KEY], "-o", color="#2E74B5")
    a1.set_title(f"cover-time cut @ short target {best['short_target']}"); a1.tick_params(axis="x", rotation=90)
    a1.set_ylabel("net_A combined total_return %"); a1.grid(alpha=0.3)
    cut2 = df[df["cover_time"] == best["cover_time"]].set_index("short_target").reindex(xorder).reset_index()
    a2.plot(range(len(xorder)), cut2[KEY], "-o", color="#C55A11")
    a2.set_xticks(range(len(xorder))); a2.set_xticklabels(xorder, rotation=90, fontsize=8)
    a2.axvline(len(xorder) - 1, color="grey", ls=":", label="no-target")
    a2.set_title(f"short-target cut @ cover {best['cover_time']}"); a2.grid(alpha=0.3); a2.legend()
    fig.tight_layout(); fig.savefig(OUTDIR / "short_1d_cuts.png", dpi=130); plt.close(fig)

    # ── console ──
    pd.set_option("display.width", 240)
    print("\n" + "=" * 100 + "\nSHORT-EXIT 2-D SWEEP (cover time × short target) — long fixed 09:25/11:59/17%\n" + "=" * 100)
    print(f"trades: {N:,} | combos: {len(df)}")
    print(f"\nBEST: cover {best['cover_time']}  short_target {best['short_target']}  ->  "
          f"net_A combined {best[KEY]:.2f}%  (win {best['net_A_win_rate_pct']}%)")
    print(f"   short leg: {best['short_leg_return_fixedbase_pct']:.2f}%  win {best['short_win_rate_pct']}%  "
          f"avg {best['avg_short_return_pct']}%  hit-target {best['pct_shorts_hitting_target']}%")
    print(f"INCUMBENT (3pm cover, no short target): net_A combined {inc[KEY]:.2f}%  "
          f"(short leg {inc['short_leg_return_fixedbase_pct']:.2f}%)")
    print(f"   best beats incumbent by {best[KEY]-inc[KEY]:+.2f} net_A pts.")
    print(f"\nTARGET vs NO-TARGET (shorts): best with-target net_A {with_t[KEY]:.2f}% "
          f"(@{with_t['cover_time']}/{with_t['short_target']}) vs best no-target {no_t[KEY]:.2f}% (@{no_t['cover_time']})"
          f"  -> {'target helps' if with_t[KEY] > no_t[KEY] else 'no-target better'} by {abs(with_t[KEY]-no_t[KEY]):.2f} pts")
    print("\npct of shorts hitting target across X (at best cover " + bc + "):")
    print(pct_tbl.to_string(index=False))
    print("\nBY-ORIGIN best short exit:")
    print(by_org.to_string(index=False))
    print("\nTOP 10:")
    print(df.head(10)[["rank", "cover_time", "short_target", KEY, "net_A_win_rate_pct",
                       "short_leg_return_fixedbase_pct", "pct_shorts_hitting_target"]].to_string(index=False))
    print(f"\nSaved -> {OUTDIR}")


if __name__ == "__main__":
    main()
