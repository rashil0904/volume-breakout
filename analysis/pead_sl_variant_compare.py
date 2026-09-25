# -*- coding: utf-8 -*-
"""pead_sl_variant_compare.py — ADDITIVE SL-level variant test. Compares the finalized SL (sl_mode=
't0_extreme': long=T0 low, short=T0 high) against a new sl_mode='t0_midpoint' (50% retracement of the T0
range: long = T0_high - 0.5*(T0_high-T0_low); short = T0_low + 0.5*(range) — i.e. the T0 candle midpoint).
ONLY the SL level changes. Fire/fill (close-basis), target (touch), exit priority, holding periods,
sizing (Rs.50,000/trade, net 0.20%) all UNCHANGED. Runs the 3 finalized books, reports per-book compare
+ flags trades whose outcome differs between the two SL levels. Does not overwrite the finalized report.
"""
import sys
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, str(Path(__file__).resolve().parent))
import pead_final_report as pfr   # reuse load(), simulate(), fy_quarter(), CAP, COST, IS_YEARS, paths

OUT = pfr.rb.RESULTS / "pead_sl_variant_compare.xlsx"
CAP, COST = pfr.CAP, pfr.COST


def sl_levels(direction, hi, lo):
    """returns (t0_extreme_sl, t0_midpoint_sl)"""
    mid = (hi + lo) / 2.0                                   # long: hi-0.5*(hi-lo); short: lo+0.5*(hi-lo) -> both = midpoint
    return (lo if direction == "long" else hi), mid


def build_pairs(a, sd, vmap, syms, direction, K, n, tgt):
    rows = []; seen = set()
    for r in a.itertuples():
        if r.symbol not in syms:
            continue
        s = sd.get(r.symbol)
        if s is None:
            continue
        dates, op, hi, lo, cl, dmap = s
        ps = dmap.get(r.ann_date); pn = int(np.searchsorted(dates, r.ann_date, "right")); pn = pn if pn < len(dates) else None
        p = pn if r.announcement_session in ("post_market", "during_market") else (ps if ps is not None else pn)
        if p is None or p < 1:
            continue
        if (dates[p] - dates[p - 1]).days > 6:                 # GUARD1
            continue
        key = (r.symbol, int(p))                                # GUARD2
        if key in seen:
            continue
        seen.add(key)
        entry = cl[p]; pv = vmap.get((r.symbol, dates[p - 1]), np.nan)
        if not (entry > 0 and pv == pv and pv > 0):
            continue
        rr = (entry - pv) / pv * 100.0
        if abs(rr) > 50.0:                                     # GUARD3
            continue
        if direction == "long" and rr < K:
            continue
        if direction == "short" and rr > -K:
            continue
        if p + n >= len(dates):
            continue
        sl_ext, sl_mid = sl_levels(direction, hi[p], lo[p])
        fo, fh, fl, fc = op[p + 1:p + 1 + n], hi[p + 1:p + 1 + n], lo[p + 1:p + 1 + n], cl[p + 1:p + 1 + n]
        out = {}
        for mode, sl in (("ext", sl_ext), ("mid", sl_mid)):
            ex, reason, dh = pfr.simulate(direction, entry, sl, n, tgt, fo, fh, fl, fc)
            gross = (ex - entry) / entry * 100.0 if direction == "long" else (entry - ex) / entry * 100.0
            net = gross - COST
            out[mode] = {"sl": round(sl, 2), "exit_date": str(dates[p + dh]), "exit_price": round(ex, 2),
                         "reason": reason, "days": dh, "net": net, "inr": net / 100.0 * CAP}
        diff = (out["ext"]["exit_date"] != out["mid"]["exit_date"]) or (abs(out["ext"]["exit_price"] - out["mid"]["exit_price"]) > 1e-6) \
            or (out["ext"]["reason"] != out["mid"]["reason"])
        rows.append({"symbol": r.symbol, "T0_date": str(dates[p]), "direction": direction, "reaction": round(rr, 2),
                     "entry": round(entry, 2), "sl_ext": out["ext"]["sl"], "sl_mid": out["mid"]["sl"],
                     "ext_exit_date": out["ext"]["exit_date"], "ext_exit_px": out["ext"]["exit_price"], "ext_reason": out["ext"]["reason"], "ext_days": out["ext"]["days"], "ext_inr": round(out["ext"]["inr"], 0),
                     "mid_exit_date": out["mid"]["exit_date"], "mid_exit_px": out["mid"]["exit_price"], "mid_reason": out["mid"]["reason"], "mid_days": out["mid"]["days"], "mid_inr": round(out["mid"]["inr"], 0),
                     "divergent": diff, "year": int(str(dates[p])[:4])})
    return pd.DataFrame(rows)


def mode_stats(df, mode):
    inr = df[f"{mode}_inr"]; reason = df[f"{mode}_reason"]; days = df[f"{mode}_days"]
    slmask = reason == "SL"
    return {"trades": len(df), "total_pnl": round(inr.sum(), 0), "win_rate": round((inr > 0).mean() * 100, 1),
            "sl_hits": int(slmask.sum()), "sl_hit_pct": round(slmask.mean() * 100, 1),
            "avg_hold": round(days.mean(), 1), "max_loss": round(inr.min(), 0),
            "avg_loss_per_sl": round(inr[slmask].mean(), 0) if slmask.any() else 0,
            "avg_per_trade": round(inr.mean(), 0)}


def main():
    a, sd, vmap = pfr.load()
    fno = set(pd.read_csv(pfr.FO_CSV, usecols=["symbol"])["symbol"].unique())
    u = pd.read_csv(pfr.UNIV); nonfo = set(u[~u["fno_eligible"]]["symbol"])
    specs = [("F&O Long", fno, "long", 4.0, 15, None), ("F&O Short", fno, "short", 4.0, 3, 8.0),
             ("Non-F&O Long", nonfo, "long", 7.0, 10, 17.0)]

    comp_rows = []; all_div = []
    pd.set_option("display.width", 240)
    for name, syms, direction, K, n, tgt in specs:
        df = build_pairs(a, sd, vmap, syms, direction, K, n, tgt)
        se, sm = mode_stats(df, "ext"), mode_stats(df, "mid")
        div = df[df.divergent].copy(); div.insert(0, "book", name); all_div.append(div)
        for mode, st in (("t0_extreme (finalized)", se), ("t0_midpoint (50% retr.)", sm)):
            comp_rows.append({"book": name, "sl_mode": mode, **st})
        print("\n" + "=" * 96 + f"\n{name}: SL t0_extreme vs t0_midpoint  ({len(df)} trades, identical event set)\n" + "=" * 96)
        print(pd.DataFrame([{"sl_mode": "t0_extreme", **se}, {"sl_mode": "t0_midpoint", **sm}]).to_string(index=False))
        print(f"  divergent trades: {int(df.divergent.sum())} ({round(df.divergent.mean()*100,1)}%)  | "
              f"SL-hits {se['sl_hits']} -> {sm['sl_hits']}  | avg loss/SL-trade Rs.{se['avg_loss_per_sl']:,.0f} -> Rs.{sm['avg_loss_per_sl']:,.0f}  | "
              f"total PnL Rs.{se['total_pnl']:,.0f} -> Rs.{sm['total_pnl']:,.0f} (delta Rs.{sm['total_pnl']-se['total_pnl']:,.0f})")

    COMP = pd.DataFrame(comp_rows); DIV = pd.concat(all_div, ignore_index=True) if all_div else pd.DataFrame()

    # ---- Excel ----
    wb = Workbook(); ws = wb.active; ws.title = "SL_Comparison"
    HF = PatternFill("solid", fgColor="1F497D"); WH = Font(bold=True, color="FFFFFF"); SUB = PatternFill("solid", fgColor="D9E1F2")
    ws.cell(1, 1, "PEAD SL-level variant: finalized T0-extreme vs 50%-retracement (T0 midpoint)").font = Font(bold=True, size=13)
    ws.cell(2, 1, "Only the SL level changes. Fire=close-basis, target=touch, priority/holds/sizing (Rs.50k, net 0.20%) unchanged. Same-day: target(intraday) precedes close-SL.").font = Font(italic=True, size=9)
    cols = ["book", "sl_mode", "trades", "total_pnl", "avg_per_trade", "win_rate", "sl_hits", "sl_hit_pct", "avg_loss_per_sl", "avg_hold", "max_loss"]
    hdr = ["Book", "SL mode", "Trades", "Total PnL ₹", "Avg/trade ₹", "Win %", "SL hits", "SL hit %", "Avg loss/SL ₹", "Avg hold", "Max loss ₹"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h); c.fill = HF; c.font = WH
    for i, row in enumerate(comp_rows, 5):
        for j, k in enumerate(cols, 1):
            cell = ws.cell(i, j, row[k]); cell.alignment = Alignment(horizontal="left" if j <= 2 else "right")
            if "PnL" in hdr[j - 1] or "loss" in hdr[j - 1].lower() or "trade ₹" in hdr[j - 1]:
                cell.number_format = '"₹"#,##0'
            if i % 4 in (1, 2):
                cell.fill = SUB if j <= 2 else PatternFill()
    for j, w in enumerate([13, 22, 8, 13, 12, 7, 8, 9, 14, 9, 13], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws2 = wb.create_sheet("Divergent_Trades")
    dcols = ["book", "symbol", "T0_date", "direction", "reaction", "entry", "sl_ext", "sl_mid",
             "ext_exit_date", "ext_exit_px", "ext_reason", "ext_days", "ext_inr",
             "mid_exit_date", "mid_exit_px", "mid_reason", "mid_days", "mid_inr"]
    for j, h in enumerate(dcols, 1):
        c = ws2.cell(1, j, h); c.fill = HF; c.font = WH
    if len(DIV):
        for i, rr in enumerate(DIV[dcols].itertuples(index=False), 2):
            for j, v in enumerate(rr, 1):
                ws2.cell(i, j, v)
    for j in range(1, len(dcols) + 1):
        ws2.column_dimensions[get_column_letter(j)].width = 12
    wb.save(OUT)

    # ---- explicit tighter-SL narrative ----
    print("\n" + "#" * 96 + "\nTIGHTER-SL EFFECT (explicit):\n" + "#" * 96)
    for name, _, _, _, _, _ in specs:
        e = COMP[(COMP.book == name) & (COMP.sl_mode.str.startswith("t0_extreme"))].iloc[0]
        m = COMP[(COMP.book == name) & (COMP.sl_mode.str.startswith("t0_midpoint"))].iloc[0]
        print(f"  {name}: SL-hit {e.sl_hit_pct}%->{m.sl_hit_pct}% ({'MORE' if m.sl_hits>e.sl_hits else 'same/less'} stops); "
              f"avg loss/SL Rs.{e.avg_loss_per_sl:,.0f}->Rs.{m.avg_loss_per_sl:,.0f} ({'smaller' if m.avg_loss_per_sl>e.avg_loss_per_sl else 'larger'} per stop); "
              f"total PnL Rs.{e.total_pnl:,.0f}->Rs.{m.total_pnl:,.0f}")
    print(f"\nSaved -> {OUT}")


if __name__ == "__main__":
    main()
